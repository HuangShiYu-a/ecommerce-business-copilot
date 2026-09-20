"""Knowledge 多格式文档生成（R1.2）：PDF 4 篇 + DOCX 3 篇。

内容遵循 knowledge-seed-enhancement-plan §2 清单 + §4.1 内容边界红线。
PDF 用 reportlab（STSong-Light 中文）；DOCX 用 python-docx（Heading 样式，供结构感知切块）。
元数据（layer/effective/scope）声明在 seed/manifest.yaml，不在此。
"""
from __future__ import annotations

from pathlib import Path

BASE = Path(__file__).resolve().parent.parent.parent / "data" / "knowledge" / "seed" / "docs"

# ---- PDF 内容（标题/节/段落）----
PDF_DOCS = {
    "KD-AS-05": {
        "title": "退货退款政策总则",
        "sections": [
            ("适用范围", "本政策适用于平台自营商品的退货退款流程。法定权利与运费规则以《七日无理由退货底线》（KD-AS-01）为准，本总则不重复规定。"),
            ("退货申请入口", "用户可通过订单详情发起退货申请；客服在售后工单中代为登记（流程见 KD-AS-06）。申请需在签收后窗口期内发起。"),
            ("审核流程", "申请经资格判定（可退类目、签收窗口、完好要求，见 KD-AS-01/02/03）后进入审核。判定结论由系统计算，不依赖人工推断。"),
            ("退款到账", "审核通过后按订单实付金额发起退款；优惠分摊与赠品扣除见 KD-AS-09。法定 7 日返还时限见 KD-AS-01。"),
            ("与既有规则的关系", "本文为流程性总则；法定规则引用 KD-AS-01，类目规则引用 KD-AS-02，完好凭证引用 KD-AS-03。"),
        ],
    },
    "KD-LG-02": {
        "title": "物流时效承诺政策 v1.0（已废止）",
        "sections": [
            ("版本说明", "本版本（v1.0）自 2026-08-31 起废止，由 v2.0（KD-LG-03）取代。依据生命周期规则（R1），本版本不再参与时效判定。"),
            ("承诺框架", "物流时效以订单承诺时效为准（KD-LG-01）；本政策定义承诺的登记、展示与异常责任框架。"),
            ("时效查询", "用户可在订单物流详情查询承诺送达时间；异常时按 KD-LG-04 处理。"),
        ],
    },
    "KD-LG-03": {
        "title": "物流时效承诺政策 v2.0（现行）",
        "sections": [
            ("版本说明", "本版本（v2.0）自 2026-09-01 起生效，取代 v1.0（KD-LG-02）。同主题版本仅当前有效版本参与判定（R1 版本链）。"),
            ("承诺框架", "物流时效以订单承诺时效为准（KD-LG-01）；承诺在订单确认时锁定，展示于订单详情。"),
            ("变更规则", "承诺一经锁定不因承运商调整而变更；不可抗力情形按异常流程处理（KD-LG-04）。"),
            ("异常责任", "超承诺未送达属物流异常，进入 KD-LG-04 处理流程；客服安抚与运营催单协同。"),
        ],
    },
    "KD-INV-02": {
        "title": "库存管理政策",
        "sections": [
            ("库存分类", "库存按动销状态分为：正常、慢销候选（判定见 KD-OP-01/OP-03）、滞销（处理见 KD-OP-03）、断货风险（KD-OP-04）。"),
            ("盘点周期", "库存账面每日与系统核对；每周对慢销与高值 SKU 抽样盘点；月度全面盘点。"),
            ("数据口径", "库存可用量 = 在库量 − 已锁定量；安全库存与再订货点见 KD-INV-01 备货原则与店铺配置。"),
            ("与备货原则的关系", "本政策管理库存状态；补货触发与确认流程见 KD-INV-01，本文不重写。"),
        ],
    },
}

# ---- DOCX 内容（标题/节/段落）----
DOCX_DOCS = {
    "KD-CS-01": {
        "title": "客服话术手册（退货/物流/售后）",
        "sections": [
            ("开场与身份核实", "客服自我介绍并核实用户身份；涉及订单操作时确认 store 权限范围。"),
            ("退货咨询话术", "用户询问能否退货时，引用资格判定结论（可退类目/窗口/完好），不承诺窗口外的退货。"),
            ("物流咨询话术", "用户询问物流时，以订单承诺时效与轨迹为准回复；异常时按 KD-LG-04 安抚并跟进。"),
            ("退款咨询话术", "退款时效遵循 KD-AS-01 法定 7 日；到账路径以实际支付方式为准，不做额外承诺。"),
            ("争议升级", "超出话术范围或涉及金额争议时，转工单升级人工复核，禁止客服口头承诺赔偿。"),
        ],
    },
    "KD-CS-02": {
        "title": "工单操作手册",
        "sections": [
            ("工单创建", "售后诉求均先建工单：登记用户、订单、诉求类型（退货/物流/退款/其他）。"),
            ("工单流转", "按类型流转：退货走 KD-AS-06 流程、物流异常走 KD-LG-04、退款走 KD-AS-07。"),
            ("工单升级", "争议、重复诉求、金额超限等升级人工复核；升级后原工单保持可跟踪。"),
            ("工单关闭", "处理完成并回访后关闭；未解决工单禁止静默关闭。"),
        ],
    },
    "KD-AS-08": {
        "title": "售后质检培训（服务评分标准）",
        "sections": [
            ("质检范围", "对售后工单按响应时效、解决率、话术规范进行抽样质检。"),
            ("评分维度", "响应及时性、信息准确（引用规则正确）、流程合规（建单/升级/关闭规范）、用户满意度。"),
            ("培训要求", "质检不达标项进入培训；规则引用错误（如错引退货窗口）为严重项，需专项纠正。"),
        ],
    },
}


def _gen_pdf(doc_id: str, spec: dict) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from reportlab.lib.styles import ParagraphStyle

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    path = BASE / "pdf" / f"{doc_id}.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    title_st = ParagraphStyle("t", fontName="STSong-Light", fontSize=16, leading=20)
    head_st = ParagraphStyle("h", fontName="STSong-Light", fontSize=12, leading=16, spaceBefore=8)
    body_st = ParagraphStyle("b", fontName="STSong-Light", fontSize=10.5, leading=15)
    story = [Paragraph(spec["title"], title_st), Spacer(1, 10)]
    for head, body in spec["sections"]:
        story.append(Paragraph(head, head_st))
        story.append(Paragraph(body, body_st))
        story.append(Spacer(1, 4))
    doc.build(story)


def _gen_docx(doc_id: str, spec: dict) -> None:
    from docx import Document

    path = BASE / "docx" / f"{doc_id}.docx"
    d = Document()
    d.add_heading(spec["title"], level=0)
    for head, body in spec["sections"]:
        d.add_heading(head, level=1)
        d.add_paragraph(body)
    d.save(str(path))


def main() -> None:
    for doc_id, spec in PDF_DOCS.items():
        _gen_pdf(doc_id, spec)
    for doc_id, spec in DOCX_DOCS.items():
        _gen_docx(doc_id, spec)
    print("generated:", len(PDF_DOCS), "PDF +", len(DOCX_DOCS), "DOCX")


if __name__ == "__main__":
    main()
